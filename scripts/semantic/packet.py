"""Coder packet for the human validation of tone scores: an Excel coding workbook and a Word codebook.

  uv run --with openpyxl --with python-docx python -m scripts.semantic.packet \
      --blind reports/semantic/validation/validation_sample_v1.csv [--mt translations.json] \
      --out reports/semantic/validation [--version v1]

The workbook has an Instructions sheet, a Codebook sheet and a Coding sheet (one row per blind sentence, 0/1 dropdowns
per dimension). Rows are put in a fixed pseudo-random order so language blocks from a top-up are mixed in. The
optional --mt file maps sample_id -> machine translation, shown in its own clearly labelled column; nothing in the
packet carries model scores or the sampling arm. Returned workbooks are scored with
`python -m scripts.semantic.run validate <file.xlsx> [...]`.
"""
from __future__ import annotations

import argparse
import csv
import json
import zipfile
import zlib
from pathlib import Path
from typing import Dict, List

from .config import DIMS, HYPOTHESES

# Operational definitions. Wording follows config.HYPOTHESES (the teacher prompts); examples are the author-labelled
# dev/hold-out sentences in reports/semantic/hypothesis_selection.md.
CODEBOOK: Dict[str, Dict] = {
    "hostility": {
        "label": "Hostility",
        "definition": "The sentence expresses hostility toward another country or actor: condemnation, accusation, "
                      "blame, contempt, or derogatory labels for a government, state, organisation or group.",
        "include": ["Condemnations and accusations, even when formulaic (\"strongly condemns\", \"baseless allegations\").",
                    "Derogatory or delegitimising labels (\"the Kiev regime\", \"neo-Nazis\", \"Zionist regime\").",
                    "Hostility voiced by a quoted speaker counts the same as the author's own."],
        "exclude": ["Neutral reporting of conflict events with no evaluative language.",
                    "A bare threat or warning with no negative evaluation of the target (code threat only).",
                    "Disagreement stated politely as policy difference."],
        "positive": ["The Kiev regime and its Western sponsors are committing crimes against civilians.",
                     "Pakistan rejects India's baseless and irresponsible allegations.",
                     "美方应立即停止向台湾出售武器，否则中方将采取坚决措施。 (The US must immediately stop selling arms to Taiwan, "
                     "otherwise China will take resolute measures.)"],
        "negative": ["The two presidents held a telephone conversation.",
                     "Tensions on the Korean Peninsula have risen sharply after a series of missile launches.",
                     "Iran will not hesitate to defend itself against any aggression with full force."],
    },
    "threat": {
        "label": "Threat",
        "definition": "The sentence is about threats or warnings of retaliation: a warning that force, punishment or "
                      "other consequences will follow, or a report of such a warning by any party.",
        "include": ["Conditional warnings (\"if X, we will respond ...\"), red lines, promises of a crushing response.",
                    "Pledges to take \"all necessary measures\" against an adversary.",
                    "Reported threats by third parties."],
        "exclude": ["Hostile description without any warning of consequences (code hostility only).",
                    "Routine statements about defence capabilities with no adversary or consequence named.",
                    "Threats from non-political sources (weather, disease, markets)."],
        "positive": ["If NATO deploys these missiles, Russia will respond with all means at its disposal.",
                     "The PLA will take all necessary measures to resolutely defeat any separatist attempt.",
                     "Any attempt to violate our sovereignty will be met with a decisive and crushing response."],
        "negative": ["Washington's reckless actions are pushing the world to the brink of nuclear war.",
                     "The Kiev regime and its Western sponsors are committing crimes against civilians.",
                     "Matvei Valkov won a gold medal in men's pole vault at the U20 Championships."],
    },
    "conciliation": {
        "label": "Conciliation",
        "definition": "The sentence calls for or describes cooperation, dialogue, partnership or friendly relations "
                      "between countries or actors.",
        "include": ["Statements of readiness to cooperate or to talk; agreements to deepen ties.",
                    "Diplomatic goodwill: thanks, congratulations, wishes of prosperity to a foreign leader or people.",
                    "Talks described as reviewing cooperation or matters of mutual interest."],
        "exclude": ["A bare procedural fact with no cooperative content (\"held a telephone conversation\").",
                    "Domestic cooperation between ministries or agencies of one state.",
                    "A ceasefire or call for restraint without a cooperative/friendly element (code deescalation only)."],
        "positive": ["The interlocutors reviewed topical matters of mutual cooperation.",
                     "He wished success, strong health, and prosperity to the Belarusian leader.",
                     "Мы готовы к диалогу с американскими партнерами на равноправной основе. (We are ready for dialogue "
                     "with our American partners on an equal footing.)"],
        "negative": ["The two presidents held a telephone conversation.",
                     "The ceasefire agreement was signed in Doha on Tuesday.",
                     "Approximately Br2.2 billion will be allocated for pension payments in September."],
    },
    "grievance": {
        "label": "Grievance",
        "definition": "The sentence complains that the speaker, its country or its people are being wronged, "
                      "victimised or treated unfairly.",
        "include": ["Complaints of injustice, illegal sanctions, double standards, ignored legitimate concerns.",
                    "Historical victimhood and demands for recognition or apology.",
                    "Rejection of accusations framed as unfair treatment of the speaker's side."],
        "exclude": ["Condemning harm done to third parties (hostility, not grievance, unless the speaker's side is the victim).",
                    "Neutral reports of damage or casualties without a complaint of unfairness.",
                    "Hostile statements that do not portray the speaker's side as wronged."],
        "positive": ["The United States has unjustly imposed illegal sanctions on our country for decades.",
                     "For decades the West has ignored our legitimate security concerns.",
                     "日本军国主义从未真正反省其侵略罪行。 (Japanese militarism has never truly reflected on its crimes of aggression.)"],
        "negative": ["The Foreign Ministry strongly condemns the terrorist attack and expresses condolences to the victims.",
                     "The Kiev regime and its Western sponsors are committing crimes against civilians.",
                     "Exports grew by 6.8 percent year on year in the first quarter."],
    },
    "escalation": {
        "label": "Escalation",
        "definition": "The sentence is about the escalation of a conflict: rising tension, intensifying hostilities, "
                      "provocations, arms races, or moves toward wider or more severe conflict.",
        "include": ["Reports that tensions rise, fighting intensifies, new strikes occur.",
                    "Accusations that an actor is provoking or pushing toward war.",
                    "Statements that a response will widen or intensify a conflict."],
        "exclude": ["A steady-state description of an ongoing conflict without intensification.",
                    "A threat that does not describe or imply widening the conflict (code threat only).",
                    "Escalation in non-conflict senses (prices, disease)."],
        "positive": ["Tensions on the Korean Peninsula have risen sharply after a series of missile launches.",
                     "Американская военщина провоцирует новый виток гонки вооружений. (The American military is "
                     "provoking a new round of the arms race.)",
                     "The situation in the Middle East continues to deteriorate rapidly, with new strikes every day."],
        "negative": ["Iran will not hesitate to defend itself against any aggression with full force.",
                     "The Kiev regime and its Western sponsors are committing crimes against civilians.",
                     "We urge the parties to avoid any steps that could further escalate the situation."],
    },
    "deescalation": {
        "label": "De-escalation",
        "definition": "The sentence urges calm, restraint or a peaceful settlement of a conflict, or reports a step "
                      "that reduces a conflict (ceasefire, truce, negotiated settlement).",
        "include": ["Calls for restraint, calm, dialogue to settle a dispute or conflict.",
                    "Ceasefires, truces, peace agreements, withdrawal or confidence-building steps.",
                    "Commitments to resolve a named dispute peacefully."],
        "exclude": ["Ceremonial goodwill or friendship with no conflict in view (code conciliation only).",
                    "Generic praise of peace as a value with no conflict or dispute referred to.",
                    "Reports that a conflict continues."],
        "positive": ["We call on all parties to exercise restraint and return to the negotiating table.",
                     "We urge the parties to avoid any steps that could further escalate the situation.",
                     "The ceasefire agreement was signed in Doha on Tuesday."],
        "negative": ["China is ready to work with Pakistan to deepen the all-weather strategic partnership.",
                     "The minister thanked his counterpart for the warm hospitality.",
                     "Tensions on the Korean Peninsula have risen sharply after a series of missile launches."],
    },
}

GENERAL_RULES = [
    "Code each sentence on its own, as written. Do not open the URL to decide a code; use it only if a sentence is an "
    "unreadable fragment, and say so in notes.",
    "Each dimension is a separate yes/no judgement: 1 = clearly present, 0 = absent. Several dimensions can be 1 at once; "
    "all can be 0 (most sentences are neutral).",
    "Quoted speech counts: code what the sentence says, whoever is speaking.",
    "If a dimension is borderline, choose the better answer and set confidence to 1 (low); explain briefly in notes.",
    "Leave the six cells blank (not 0) only if you cannot read the sentence at all; write 'cannot read' in notes. "
    "Blank cells are excluded from scoring; 0 means 'absent'.",
    "Headlines, datelines, lists and boilerplate are coded like any sentence (usually all 0).",
    "Machine translation (where given) is an aid only, produced automatically and possibly wrong. Code the original "
    "text; if you rely on the translation alone, note 'MT only'.",
]

INSTRUCTIONS = [
    ("Purpose", "You are providing human ground truth for six tone dimensions that a computer model assigns to sentences "
                "from official statements and state media. Your codes are compared with the model afterwards."),
    ("Blind protocol", "The sheet shows no model scores and nothing about how sentences were selected. Do not look for "
                       "the model output, do not discuss individual sentences with the other coder before both of you "
                       "have finished, and do not use any AI tool or chatbot to suggest codes. Translation tools for "
                       "understanding a sentence are fine; note it in notes."),
    ("Two independent coders", "Ideally two people code the full sheet independently, each in their own copy of this "
                               "file. Agreement between coders (Krippendorff's alpha) is reported per dimension."),
    ("Before you start", "Read the Codebook sheet (or validation_codebook_v1.docx). Code the first 20 rows, re-read the "
                         "codebook, then continue. Put your name or initials in the 'coder' column (every row you code)."),
    ("Columns", "Fill only: the six dimension columns (0/1 dropdowns), coder, confidence (1 = low, 2 = medium, "
                "3 = high; your overall confidence for the row), english_gloss (optional, your own short gloss for "
                "non-English rows) and notes. Do not edit, sort-and-save, delete or insert rows in the other columns; "
                "sample_id must stay intact."),
    ("Languages", "If you cannot read a language and the machine translation is not enough, leave the six cells blank "
                  "and write 'cannot read' in notes. Partial coding is fine; blank rows are skipped."),
    ("Returning the file", "Save as .xlsx under a new name, e.g. validation_coding_v1_<initials>.xlsx, and send it back "
                           "to Jonathan Walberg. Do not convert to Google Sheets and back unless the dropdown values "
                           "remain plain 0/1."),
    ("Time", "Roughly 30 to 45 seconds per sentence; the full sheet takes about 4 to 5 hours. Breaks are fine."),
]

CODING_COLS = ["sample_id", "country", "source", "lang", "date", "text", "machine_translation", "english_gloss", "url",
               *DIMS, "coder", "confidence", "notes"]


def order_rows(rows: List[Dict[str, str]]) -> List[Dict[str, str]]:
    """Fixed pseudo-random order by sample_id (CRC), so top-up rows are mixed in and every copy is identical."""
    return sorted(rows, key=lambda r: (zlib.crc32(("packet:" + r["sample_id"]).encode()), r["sample_id"]))


def build_xlsx(rows: List[Dict[str, str]], mt: Dict[str, str], path: Path, version: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    bold, wrap = Font(bold=True), Alignment(wrap_text=True, vertical="top")
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    code_fill = PatternFill("solid", fgColor="FFF6D5")

    ins = wb.active
    ins.title = "Instructions"
    ins.append([f"Tone validation coding sheet {version}", ""])
    ins["A1"].font = Font(bold=True, size=14)
    ins.append(["", ""])
    for k, v in INSTRUCTIONS:
        ins.append([k, v])
    ins.append(["", ""])
    ins.append(["General coding rules", ""])
    ins.cell(ins.max_row, 1).font = bold
    for i, rule in enumerate(GENERAL_RULES, 1):
        ins.append([f"Rule {i}", rule])
    ins.column_dimensions["A"].width = 24
    ins.column_dimensions["B"].width = 110
    for row in ins.iter_rows(min_row=3):
        row[0].font = bold
        for c in row:
            c.alignment = wrap

    cb = wb.create_sheet("Codebook")
    cb.append(["dimension", "operational definition", "includes", "excludes", "code 1 (examples)", "code 0 (examples)",
               "model prompt (for reference)"])
    for d in DIMS:
        e = CODEBOOK[d]
        cb.append([e["label"], e["definition"], "\n".join("- " + x for x in e["include"]),
                   "\n".join("- " + x for x in e["exclude"]), "\n".join("- " + x for x in e["positive"]),
                   "\n".join("- " + x for x in e["negative"]), HYPOTHESES[d]])
    for col, w in zip("ABCDEFG", (14, 40, 45, 45, 50, 50, 34)):
        cb.column_dimensions[col].width = w
    for row in cb.iter_rows():
        for c in row:
            c.alignment = wrap
    for c in cb[1]:
        c.font, c.fill = bold, head_fill
    cb.freeze_panes = "B2"

    ws = wb.create_sheet("Coding")
    ws.append(CODING_COLS)
    for r in rows:
        ws.append([r["sample_id"], r["country"], r["source"], r["lang"], r["date"], r["text"],
                   mt.get(r["sample_id"], ""), "", r["url"], *[""] * len(DIMS), "", "", ""])
    n = len(rows) + 1
    widths = {"sample_id": 9, "country": 8, "source": 16, "lang": 6, "date": 11, "text": 70, "machine_translation": 55,
              "english_gloss": 30, "url": 22, "coder": 10, "confidence": 11, "notes": 30, **{d: 9 for d in DIMS}}
    letter = {c: ws.cell(1, i + 1).column_letter for i, c in enumerate(CODING_COLS)}
    for c, w in widths.items():
        ws.column_dimensions[letter[c]].width = w
    for row in ws.iter_rows(min_row=1, max_row=n):
        for c in row:
            c.alignment = wrap
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    ws.cell(1, CODING_COLS.index("machine_translation") + 1).value = "machine_translation (AI-generated, may be wrong)"
    dv = DataValidation(type="list", formula1='"0,1"', allow_blank=True, showErrorMessage=True,
                        errorTitle="0 or 1", error="Enter 0 (absent) or 1 (present), or leave blank if unreadable.")
    conf = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True, showErrorMessage=True,
                          errorTitle="1-3", error="Confidence: 1 = low, 2 = medium, 3 = high.")
    ws.add_data_validation(dv)
    ws.add_data_validation(conf)
    dv.add(f"{letter[DIMS[0]]}2:{letter[DIMS[-1]]}{n}")
    conf.add(f"{letter['confidence']}2:{letter['confidence']}{n}")
    for d in DIMS:
        for (c,) in ws.iter_rows(min_row=2, max_row=n, min_col=CODING_COLS.index(d) + 1, max_col=CODING_COLS.index(d) + 1):
            c.fill = code_fill
            c.alignment = Alignment(horizontal="center", vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{letter['notes']}{n}"
    wb.active = wb.index(ins)
    wb.save(path)


def build_docx(path: Path, version: str) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    # East Asian fallback so the Chinese examples render (Calibri has no CJK glyphs).
    from docx.oxml.ns import qn
    doc.styles["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    doc.add_heading(f"Tone validation codebook ({version})", 0)
    doc.add_paragraph("Human coding of six tone dimensions in sentences from official statements and state media. "
                      "Each dimension is coded 0 (absent) or 1 (clearly present) for each sentence, independently of "
                      "the others. The definitions follow the prompts the scoring model was built on; the examples are "
                      "the hand-labelled sentences used when those prompts were chosen.")
    doc.add_heading("Procedure", 1)
    for k, v in INSTRUCTIONS:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(k + ". ").bold = True
        p.add_run(v)
    doc.add_heading("General coding rules", 1)
    for rule in GENERAL_RULES:
        doc.add_paragraph(rule, style="List Number")
    for d in DIMS:
        e = CODEBOOK[d]
        doc.add_heading(f"{e['label']} ({d})", 1)
        doc.add_paragraph(e["definition"])
        for title, key in (("Includes", "include"), ("Excludes", "exclude"), ("Code 1, examples", "positive"),
                           ("Code 0, examples", "negative")):
            doc.add_heading(title, 2)
            for x in e[key]:
                doc.add_paragraph(x, style="List Bullet")
        p = doc.add_paragraph()
        p.add_run("Model prompt (reference only): ").italic = True
        p.add_run(HYPOTHESES[d]).italic = True
    doc.save(path)
    content_types_first(path)


def content_types_first(path: Path) -> None:
    """Rewrite a .docx/.xlsx so [Content_Types].xml is the first zip entry (Word may otherwise call it damaged)."""
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    if items and items[0][0].filename == "[Content_Types].xml":
        return
    items.sort(key=lambda x: x[0].filename != "[Content_Types].xml")
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            z.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED)
    tmp.replace(path)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--blind", required=True, help="blind sample CSV (no scores)")
    ap.add_argument("--mt", default=None, help="JSON {sample_id: machine translation}")
    ap.add_argument("--out", required=True)
    ap.add_argument("--version", default="v1")
    a = ap.parse_args()
    rows = order_rows(list(csv.DictReader(open(a.blind, encoding="utf-8"))))
    mt = json.loads(Path(a.mt).read_text(encoding="utf-8")) if a.mt else {}
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    xl = out / f"validation_coding_{a.version}.xlsx"
    build_xlsx(rows, mt, xl, a.version)
    content_types_first(xl)
    build_docx(out / f"validation_codebook_{a.version}.docx", a.version)
    print(json.dumps({"rows": len(rows), "with_mt": sum(r["sample_id"] in mt for r in rows), "xlsx": str(xl)}))


if __name__ == "__main__":
    main()
